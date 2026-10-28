"""Pruebas de actualización diaria; variantes futuras son sintéticas y temporales."""
import copy
import json
import shutil
import tempfile
import unittest
from datetime import datetime,timedelta
from pathlib import Path
from openpyxl import load_workbook
from actualizar_bna import ROOT,TZ
from actualizar_oficial import apply_snapshot,BOOK

class Daily(unittest.TestCase):
    def test_same_file_append_replace_keep_history_and_errors(self):
        snap=json.loads((ROOT/'evidencia/inicio_oficial/fuente.json').read_text())
        captured=datetime.fromisoformat(snap['captured_at'].replace('Z','+00:00'))
        with tempfile.TemporaryDirectory(dir=ROOT) as folder:
            book=Path(folder)/BOOK.name
            evidence=Path(folder)/'fuente.json'
            shutil.copy2(BOOK,book)
            evidence.write_text(json.dumps(snap))
            result=apply_snapshot(book,snap,evidence,now=captured+timedelta(seconds=1))
            self.assertEqual(result['filas'],27)
            self.assertTrue(Path(result['respaldo']).exists())
            tomorrow=copy.deepcopy(snap)
            then=captured+timedelta(days=1)
            tomorrow['captured_at']=then.isoformat()
            tomorrow['queried_to']=then.astimezone(TZ).date().isoformat()
            evidence.write_text(json.dumps(tomorrow))
            result=apply_snapshot(book,tomorrow,evidence,now=then+timedelta(seconds=1))
            self.assertEqual(result['filas'],28)
            self.assertFalse(result['dato_hoy'])
            wb=load_workbook(book,data_only=True)
            self.assertEqual(wb['Cotizaciones']['A34'].value.day,28)
            self.assertEqual(wb['Cotizaciones']['F34'].value.day,27)
            self.assertIn('SIN FILA',wb['Cotizaciones']['G34'].value)
            self.assertEqual(wb['Cotizaciones']['B7'].value,1495)
            wb.close()
            # Dato nuevo y una corrección, exclusivamente para comprobar el mecanismo.
            tomorrow['rows'].append(['28/10/2026','1600,0000','1650,0000','1610,0000','1660,0000'])
            tomorrow['captured_at']=(then+timedelta(minutes=1)).isoformat()
            evidence.write_text(json.dumps(tomorrow))
            apply_snapshot(book,tomorrow,evidence,now=then+timedelta(minutes=2))
            wb=load_workbook(book,data_only=True)
            self.assertEqual(wb['Cotizaciones'].max_row,34)
            self.assertEqual(wb['Cotizaciones']['B34'].value,1600)
            self.assertEqual(wb['Cotizaciones']['F34'].value.day,28)
            wb.close()
            before=book.read_bytes()
            with self.assertRaises(ValueError):
                apply_snapshot(book,snap,evidence,now=then)
            self.assertEqual(before,book.read_bytes())
            (book.parent/('~$'+book.name)).touch()
            with self.assertRaises(ValueError):
                apply_snapshot(book,tomorrow,evidence,now=then+timedelta(minutes=2))
            self.assertEqual(before,book.read_bytes())

if __name__=='__main__': unittest.main()
