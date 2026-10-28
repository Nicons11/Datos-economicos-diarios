"""Casos de integridad y migración en copias temporales; sin publicar datos de prueba."""
import copy
import json
import shutil
import tempfile
import unittest
from datetime import date, datetime
from pathlib import Path
from openpyxl import load_workbook
from fuentes_economicas import TZ, parse_sofr, parse_rofex, parse_index, parse_cronista, month_offset
from actualizar_economicos import apply_results, existing_rows, ROOT, BOOK

TODAY = date(2026,10,28)
NOW = datetime(2026,10,28,19,5,tzinfo=TZ)

class Sources(unittest.TestCase):
    def test_rofex_uses_settlement_zero_volume_and_thirteen_months(self):
        p = {'data':[dict(dateTime='2026-10-28T00:00:00Z',symbol='DLR102027',
                          product='DLR',settlement=1900,close=0,volume=0)]}
        rows = parse_rofex(p,TODAY)
        self.assertEqual(len(rows),13)
        self.assertEqual(rows[-1][:4],[TODAY,'DLR102027',date(2027,10,1),1900])
        self.assertIsNone(rows[0][3]); self.assertIn('SIN AJUSTE',rows[0][4])
        self.assertEqual(month_offset(date(2026,12,31),12),date(2027,12,1))
        p['data'][0]['settlement'] = 0
        with self.assertRaises(ValueError): parse_rofex(p,TODAY)

    def test_sofr_rate_effective_date_and_revisions(self):
        p = {'refRates':[dict(effectiveDate='2026-10-27',type='SOFR',percentRate=3.95,revisionIndicator='*')]}
        r = parse_sofr(p,TODAY)[0]
        self.assertEqual(r[0],date(2026,10,27)); self.assertAlmostEqual(r[1],0.0395)
        self.assertEqual(r[2],'REVISADO')
        p['refRates'] *= 2
        with self.assertRaises(ValueError): parse_sofr(p,TODAY)

    def test_investing_dates_numbers_ranges_and_stale(self):
        html = '<h1>Bovespa Historical Data</h1><table><tr>' + ''.join('<th>'+s+'</th>' for s in
                ['Date','Price','Open','High','Low','Vol.','Change %']) + '</tr><tr>' + ''.join('<td>'+s+'</td>' for s in
                ['Oct 28, 2026','170,000','169,000','171,000','168,000','12M','+1.20%']) + '</tr></table>'
        r = parse_index(html,TODAY,'Bovespa')[0]
        self.assertEqual(r[1],170000); self.assertAlmostEqual(r[5],0.012)
        self.assertIn('PROVISIONAL',r[-1])
        for invalid in [html.replace('170,000','172,000'),html.replace('Oct 28','Oct 08'),html.replace('Date','Fecha')]:
            with self.assertRaises(ValueError): parse_index(invalid,TODAY,'Bovespa')

    def test_cronista_exact_identity_and_spread(self):
        ms = int(NOW.timestamp()*1000)
        item = dict(Ric='ARS=BNAR',Nombre='Dólar BNA',Compra=1510,Venta=1560,
                    UltimaActualizacion=f'/Date({ms})/',historico=[])
        def html(x): return 'Fusion.contentCache='+json.dumps({'markets-general':{'query':{'data':[x]}}})+';'
        self.assertEqual(parse_cronista(html(item),TODAY)[0][1:3],[1510,1560])
        for field,value in [('Ric','ARSB'),('Compra',1600),('Venta',None)]:
            invalid = copy.deepcopy(item); invalid[field] = value
            with self.assertRaises(ValueError): parse_cronista(html(invalid),TODAY)

class Workbook(unittest.TestCase):
    def test_migration_partial_update_backup_and_repeat(self):
        with tempfile.TemporaryDirectory() as folder:
            book = Path(folder)/BOOK.name; shutil.copy2(BOOK,book)
            original = book.read_bytes()
            before = load_workbook(book); history = list(before['Cotizaciones'].values); before.close()
            result = {'Cotizaciones':[[TODAY,1510,1560,'OBSERVADO']],
                      'SOFR':[[date(2026,10,27),0.0395,'PUBLICADO']]}
            report = apply_results(book,result,{'Bovespa':'HTTP 403'},NOW)
            self.assertEqual(Path(report['respaldo']).read_bytes(),original)
            wb = load_workbook(book)
            self.assertEqual(list(wb['Histórico Errepar'].values),history)
            self.assertEqual(wb['Fuente'].max_row,5)
            self.assertEqual(wb['Cotizaciones']['B7'].value,1510)
            self.assertEqual(wb['SOFR']['B7'].number_format,'0.00%')
            self.assertEqual(len(existing_rows(wb,'SOFR')),1)
            wb.close()
            apply_results(book,{'SOFR':[[date(2026,10,27),0.0396,'REVISADO']]},{'Cotizaciones':'bloqueada'},NOW)
            wb = load_workbook(book)
            self.assertEqual(wb['Cotizaciones']['B7'].value,1510)
            self.assertEqual(wb['SOFR'].max_row,7)
            self.assertEqual(wb['SOFR']['B7'].value,0.0396)
            wb.close()
            stable = book.read_bytes()
            (book.parent/('~$'+book.name)).touch()
            with self.assertRaises(ValueError): apply_results(book,result,{},NOW)
            self.assertEqual(book.read_bytes(),stable)
            with self.assertRaises(ValueError): apply_results(book,{}, {'SOFR':'error'},NOW)
            self.assertEqual(book.read_bytes(),stable)

if __name__ == '__main__': unittest.main()
