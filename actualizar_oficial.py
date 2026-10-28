#!/usr/bin/env python3
"""Actualiza siempre el mismo Excel con calendario diario y las cuatro series BNA."""
import argparse
import bisect
import hashlib
import json
import shutil
import sys
import tempfile
from datetime import date, datetime, timedelta, timezone
from decimal import Decimal
from pathlib import Path
from actualizar_bna import ROOT, TZ, HEADERS, URL, parse_date, validate_snapshot, fetch_live, local_path

BOOK = ROOT / 'Tipo_de_Cambio_BNA.xlsx'
OUTPUT_HEADERS = [*HEADERS, 'Fecha efectiva', 'Estado', 'Captura (Argentina)']

def apply_snapshot(book, snap, evidence, now=None):
    from openpyxl import load_workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.formatting.rule import FormulaRule
    now = now or datetime.now(TZ)
    today = now.astimezone(TZ).date()
    records, captured = validate_snapshot(snap)
    if captured.astimezone(TZ).date() != today or date.fromisoformat(snap['queried_to']) != today:
        raise ValueError('Se requiere una consulta nueva que cubra hasta hoy; no reutilizar evidencia histórica')
    if captured > now + timedelta(minutes=5):
        raise ValueError('La captura tiene una hora futura')
    if (book.parent / ('~$' + book.name)).exists():
        raise ValueError('Excel está abierto: cerrar el archivo para evitar sobrescribir cambios')
    wb = load_workbook(book)
    out, src = wb['Cotizaciones'], wb['Fuente']
    if [out.cell(6,c).value for c in range(1,9)] != OUTPUT_HEADERS:
        raise ValueError('Cambió la estructura del archivo principal')
    if [src.cell(6,c).value for c in range(1,6)] != HEADERS:
        raise ValueError('Cambió la estructura de Fuente')
    history = {}
    for row in src.iter_rows(min_row=7, values_only=True):
        if row[0] is None:
            continue
        d = parse_date(row[0])
        if d in history:
            raise ValueError('Fechas duplicadas en Fuente del Excel')
        vals = tuple(Decimal(str(x)) for x in row[1:5])
        if any(not n.is_finite() or n <= 0 for n in vals):
            raise ValueError('Valor inválido en el historial del Excel')
        old_capture = row[5].replace(tzinfo=timezone.utc) if isinstance(row[5],datetime) else datetime.fromisoformat(str(row[5]))
        if old_capture.tzinfo is None:
            raise ValueError('Fecha de captura histórica sin zona horaria')
        history[d] = (vals, old_capture)
    for d, rates in records.items():
        if d in history and history[d][1] > captured:
            raise ValueError('La captura es anterior a los datos ya guardados')
        history[d] = (rates, captured)
    dates = sorted(history)
    start = dates[0]
    if start > today:
        raise ValueError('Historial sin antecedente')
    rows = []
    d = start
    while d <= today:
        effective = dates[bisect.bisect_right(dates,d)-1]
        vals, stamp = history[effective]
        status = 'PUBLICADO EN ERREPAR' if d == effective else 'SIN FILA PARA FECHA: anterior disponible'
        if effective == stamp.astimezone(TZ).date():
            status += '; PROVISIONAL: verificar cierre'
        rows.append([d,*[float(x) for x in vals],effective,status,stamp.astimezone(TZ).replace(tzinfo=None)])
        d += timedelta(days=1)
    # Un solo archivo principal; copia de respaldo antes del reemplazo atómico.
    for sheet in (out,src):
        if sheet.max_row >= 7:
            sheet.delete_rows(7,sheet.max_row-6)
    for row in rows:
        out.append(row)
    for d in dates:
        vals, stamp = history[d]
        src.append([d,*[float(x) for x in vals],stamp.astimezone(timezone.utc).replace(tzinfo=None)])
    stamp_text=captured.astimezone(TZ).strftime('%d/%m/%Y %H:%M:%S')
    out['A3']=f'Última consulta: {stamp_text} (Argentina). Actualización diaria programada a las 19:05.'
    if today in records:
        out['A4']='Dato de hoy presente. Las cotizaciones de la fecha de captura se señalan como provisionales.'
    else:
        out['A4']='Errepar todavía no muestra fila para hoy. Se usa anterior disponible con su fecha efectiva.'
    src['A2']=f'Fuente: {URL}'
    src['A3']=f'Consulta: {stamp_text} (Argentina). Intervalo {snap["queried_from"]} a {snap["queried_to"]}.'
    src['A4']=f'Evidencia: {evidence.relative_to(ROOT)}'
    src['A5']=f'SHA256: {hashlib.sha256(evidence.read_bytes()).hexdigest()}'
    for sheet in (out,src):
        for row in sheet.iter_rows(min_row=7):
            for cell in row:
                cell.font=Font(name='Arial',size=10,color='172B3A')
                cell.alignment=Alignment(vertical='center')
            row[0].number_format='dd/mm/yyyy'
            for cell in row[1:5]:
                cell.number_format='#,##0.0000'
            if sheet == src:
                row[5].number_format='dd/mm/yyyy hh:mm:ss'
        sheet.auto_filter.ref=f'A6:{"H" if sheet==out else "F"}{sheet.max_row}'
    for row in out.iter_rows(min_row=7):
        row[5].number_format='dd/mm/yyyy'
        row[7].number_format='dd/mm/yyyy hh:mm:ss'
        row[6].alignment=Alignment(wrap_text=True,vertical='center')
        out.row_dimensions[row[0].row].height=30
    out.conditional_formatting._cf_rules.clear()
    out.conditional_formatting.add(f'G7:G{out.max_row}',FormulaRule(formula=['OR(ISNUMBER(SEARCH("PROVISIONAL",G7)),ISNUMBER(SEARCH("SIN FILA",G7)))'],fill=PatternFill('solid',fgColor='FFF0CD')))
    backupdir=book.parent/'respaldos'
    backupdir.mkdir(exist_ok=True)
    stamp=now.strftime('%Y%m%d_%H%M%S_%f')
    backup=backupdir/f'{book.stem}_{stamp}.xlsx'
    with tempfile.NamedTemporaryFile(dir=book.parent,suffix='.xlsx',delete=False) as f:
        temporary=Path(f.name)
    try:
        wb.save(temporary)
        check=load_workbook(temporary,read_only=True,data_only=True)
        if check['Cotizaciones'].max_row != 6+len(rows):
            raise ValueError('No se pudo verificar el archivo guardado')
        check.close()
        if (book.parent/('~$'+book.name)).exists():
            raise ValueError('Excel se abrió durante la actualización; se mantiene el archivo anterior')
        shutil.copy2(book,backup)
        temporary.replace(book)
    finally:
        temporary.unlink(missing_ok=True)
    return {'archivo':str(book),'filas':len(rows),'hasta':today.isoformat(),'dato_hoy':today in records,'consulta':stamp_text,'respaldo':str(backup)}

def main():
    p=argparse.ArgumentParser(description=__doc__)
    p.add_argument('--snapshot',help='JSON NUEVO obtenido del navegador en la misma ejecución')
    args=p.parse_args()
    if args.snapshot:
        evidence=local_path(args.snapshot)
        snap=json.loads(evidence.read_text())
    else:
        from openpyxl import load_workbook
        wb=load_workbook(BOOK,read_only=True,data_only=True)
        olddates=[parse_date(r[0]) for r in wb['Fuente'].iter_rows(min_row=7,values_only=True) if r[0] is not None]
        wb.close()
        today=datetime.now(TZ).date()
        first=max(min(olddates), min(max(olddates)-timedelta(days=7),today-timedelta(days=7)))
        snap,evidence=fetch_live(first,today,ROOT/'evidencia'/datetime.now(TZ).strftime('%Y%m%d_%H%M%S_%f'))
    print(json.dumps(apply_snapshot(BOOK,snap,evidence),ensure_ascii=False))

if __name__=='__main__':
    try:
        main()
    except Exception as exc:
        print(f'ERROR: {exc}. Se conserva el Excel anterior.',file=sys.stderr)
        sys.exit(1)
