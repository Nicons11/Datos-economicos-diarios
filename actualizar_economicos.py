"""Actualiza el mismo Excel de Drive con BNA, índices, futuros DLR y SOFR."""
import argparse
import json
import shutil
import tempfile
from datetime import date, datetime, timedelta
from pathlib import Path
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from actualizar_oficial import BOOK
from fuentes_economicas import TZ, CRONISTA, SOFR, ROFEX, INDICES, cronista, investing, rofex, sofr

ROOT = Path(__file__).resolve().parent
REPORT = ROOT / 'evidencia' / 'resultado.json'
SCHEMAS = {
    'Cotizaciones': ['Fecha', 'BNA compra (ARS/USD)', 'BNA venta (ARS/USD)', 'Estado', 'Captura (Argentina)'],
    **{name: ['Fecha', 'Cierre (puntos)', 'Apertura (puntos)', 'Máximo (puntos)', 'Mínimo (puntos)',
              'Variación diaria', 'Estado', 'Captura (Argentina)'] for name in INDICES},
    'ROFEX': ['Fecha rueda', 'Contrato', 'Mes contrato', 'Ajuste (ARS/USD)', 'Estado', 'Captura (Argentina)'],
    'SOFR': ['Fecha efectiva', 'SOFR anual (%)', 'Estado', 'Captura (Argentina)'],
}
URLS = {'Cotizaciones': CRONISTA, 'ROFEX': ROFEX, 'SOFR': SOFR,
        **{name: f'https://www.investing.com/indices/{slug}-historical-data' for name,(slug,_) in INDICES.items()}}

def as_date(value):
    return value.date() if isinstance(value, datetime) else value

def existing_rows(wb, name):
    if name not in wb.sheetnames: return []
    ws = wb[name]
    headers = [ws.cell(6, c).value for c in range(1, len(SCHEMAS[name]) + 1)]
    if headers != SCHEMAS[name]:
        if name == 'Cotizaciones' and headers[:3] == ['Fecha', 'Billete Compra', 'Billete Venta']:
            return []
        raise ValueError(f'Cambió la estructura de {name}; revisar antes de sobrescribir')
    rows = []
    for row in ws.iter_rows(min_row=7, max_col=len(headers), values_only=True):
        if row[0] is None: continue
        row = list(row); row[0] = as_date(row[0])
        if not isinstance(row[0], date): raise ValueError(f'Fecha inválida en {name}')
        if name == 'ROFEX': row[2] = as_date(row[2])
        rows.append(row)
    keys = [(r[0], r[1]) if name == 'ROFEX' else r[0] for r in rows]
    if len(set(keys)) != len(keys): raise ValueError(f'Historial duplicado en {name}')
    return rows

def collect(book, evidence, now):
    wb = load_workbook(book, read_only=True, data_only=True)
    today = now.date(); results = {}; errors = {}
    for name in SCHEMAS:
        try:
            old = existing_rows(wb, name)
            start = max(r[0] for r in old) - timedelta(days=7) if old else today - timedelta(days=31)
            if name == 'Cotizaciones': rows = cronista(evidence, today)
            elif name == 'SOFR': rows = sofr(evidence, today, start)
            elif name == 'ROFEX': rows = rofex(evidence, today, start)
            else: rows = investing(name, evidence, today)
            # Las páginas históricas muestran un mes: nunca ocultar huecos tras una inactividad larga.
            if old and name in ('Cotizaciones', *INDICES) and min(r[0] for r in rows) > max(r[0] for r in old):
                raise ValueError('El historial visible no solapa con el Excel: recuperar el intervalo faltante')
            results[name] = rows
        except Exception as exc:
            errors[name] = f'{type(exc).__name__}: {exc}'
    wb.close()
    return results, errors

def format_sheet(ws, name, rows, now):
    headers = SCHEMAS[name]
    ws['A1'] = {'Cotizaciones': 'Tipo de cambio nominal · Dólar BNA', 'ROFEX': 'Futuros de dólar · Ajustes A3',
                'SOFR': 'Secured Overnight Financing Rate · Overnight'}.get(name, INDICES.get(name, ('', name))[1])
    ws['A2'] = 'Fuente: ' + URLS[name]
    ws['A3'] = 'Última consulta correcta: ' + now.strftime('%d/%m/%Y %H:%M:%S') + ' (Argentina)'
    notes = {'Cotizaciones': 'Compra y venta BNA de El Cronista. Fechas publicadas; no se inventan cotizaciones para feriados.',
             'ROFEX': 'Precio de ajuste oficial, incluso sin volumen. Mes actual y 12 siguientes por rueda; faltantes quedan vacíos.',
             'SOFR': 'Tasa anual overnight. Fecha efectiva del dato; no es la fecha de publicación ni Term SOFR.',
             'Merval': 'MSCI Merval según Investing. Serie separada del antiguo S&P Merval; no se empalman bases.'}
    ws['A4'] = notes.get(name, 'Índice al cierre de cada fecha publicada. La fila del día queda provisional hasta la próxima consulta.')
    ws['A5'] = 'Última fecha disponible: ' + max(r[0] for r in rows).strftime('%d/%m/%Y')
    for r in range(1,6):
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=len(headers))
        ws.cell(r,1).alignment = Alignment(wrap_text=True, vertical='center')
        ws.row_dimensions[r].height = 30 if r in (2,4) else 23
    ws['A1'].font = Font(name='Arial', size=16, bold=True, color='172B3A')
    for c, header in enumerate(headers, 1):
        cell = ws.cell(6,c,header); cell.font = Font(name='Arial',bold=True,color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='172B3A'); cell.alignment = Alignment(wrap_text=True)
        from openpyxl.utils import get_column_letter
        ws.column_dimensions[get_column_letter(c)].width = 23
    ws.row_dimensions[6].height = 32
    for row in rows: ws.append(row)
    for row in ws.iter_rows(min_row=7):
        for cell in row:
            cell.font = Font(name='Arial',size=10,color='172B3A')
            cell.alignment = Alignment(vertical='center')
            if isinstance(cell.value, (date,datetime)): cell.number_format = 'dd/mm/yyyy'
            elif isinstance(cell.value, (float,int)): cell.number_format = '#,##0.00'
        row[-1].number_format = 'dd/mm/yyyy hh:mm:ss'
        if name in INDICES: row[5].number_format = '0.00%'
        if name == 'SOFR': row[1].number_format = '0.00%'
        if name == 'ROFEX': row[2].number_format = 'mmm yyyy'
    state_column = len(headers) - 1
    ws.column_dimensions[get_column_letter(state_column)].width = 58
    ws.column_dimensions[get_column_letter(len(headers))].width = 27
    ws.freeze_panes = 'B7'; ws.auto_filter.ref = f'A6:{get_column_letter(len(headers))}{ws.max_row}'
    ws.sheet_view.showGridLines = False

def apply_results(book, results, errors, now):
    if not results: raise ValueError('Ninguna fuente validó; se conserva el libro anterior')
    lock = book.parent / ('~$' + book.name)
    if lock.exists(): raise ValueError('Excel abierto; cerrar antes de actualizar')
    wb = load_workbook(book)
    for name, new in results.items():
        old = existing_rows(wb, name)
        if name == 'Cotizaciones' and name in wb.sheetnames and not old:
            if 'Histórico Errepar' in wb.sheetnames: raise ValueError('Archivo histórico ya existe: revisar migración')
            wb[name].title = 'Histórico Errepar'
        key = (lambda r: (r[0],r[1])) if name == 'ROFEX' else (lambda r: r[0])
        merged = {key(r): r for r in old}
        stamp = now.replace(tzinfo=None)
        for r in new: merged[key(r)] = [*r, stamp]
        rows = [merged[k] for k in sorted(merged)]
        position = wb.sheetnames.index(name) if name in wb.sheetnames else len(wb.sheetnames)
        if name in wb.sheetnames: del wb[name]
        ws = wb.create_sheet(name, position)
        format_sheet(ws, name, rows, now)
    # Preferencia del usuario: Fuente tiene exactamente cinco renglones de metadatos.
    if 'Fuente' in wb.sheetnames: del wb['Fuente']
    src = wb.create_sheet('Fuente')
    src.append(['Fuentes · ' + now.strftime('%d/%m/%Y %H:%M:%S') + ' Argentina', 'Estado de la última consulta'])
    for group, names in [('BNA · El Cronista',['Cotizaciones']), ('Índices · Investing',list(INDICES)),
                         ('ROFEX · A3',['ROFEX']), ('SOFR · New York Fed',['SOFR'])]:
        status = '; '.join(name + ': ' + (errors[name] if name in errors else 'actualizado') for name in names)
        src.append([group,status])
    src.column_dimensions['A'].width = 35; src.column_dimensions['B'].width = 115
    for row in src:
        for cell in row: cell.alignment = Alignment(wrap_text=True,vertical='center')
        src.row_dimensions[row[0].row].height = 45
    backupdir = book.parent / 'respaldos'; backupdir.mkdir(exist_ok=True)
    backup = backupdir / f'{book.stem}_{now:%Y%m%d_%H%M%S_%f}.xlsx'
    with tempfile.NamedTemporaryFile(dir=book.parent, suffix='.xlsx', delete=False) as handle:
        tmp = Path(handle.name)
    try:
        wb.save(tmp); wb.close()
        check = load_workbook(tmp, read_only=True, data_only=True)
        if check['Fuente'].max_row != 5: raise ValueError('No se pudo validar Fuente')
        for name in results:
            if check[name].max_row < 7: raise ValueError('Hoja de datos vacía')
        check.close()
        if lock.exists(): raise ValueError('Excel se abrió durante la consulta')
        shutil.copy2(book, backup); tmp.replace(book)
    finally:
        tmp.unlink(missing_ok=True)
    return {'actualizadas': list(results), 'errores': errors, 'respaldo': str(backup), 'consulta': now.isoformat()}

def main():
    p = argparse.ArgumentParser(description=__doc__); p.add_argument('--book', type=Path, default=BOOK)
    args = p.parse_args(); now = datetime.now(TZ)
    if (args.book.parent / ('~$' + args.book.name)).exists(): raise ValueError('Excel abierto; cerrar antes de actualizar')
    evidence = ROOT / 'evidencia' / now.strftime('%Y%m%d_%H%M%S_%f'); evidence.mkdir(parents=True)
    results, errors = collect(args.book, evidence, now)
    report = apply_results(args.book, results, errors, now)
    REPORT.write_text(json.dumps(report, ensure_ascii=False, indent=2))
    (evidence / 'resultado.json').write_text(REPORT.read_text())
    print(json.dumps(report, ensure_ascii=False))

def check_result():
    # Llamar DESPUÉS de subir las fuentes correctas: Actions avisa la actualización parcial.
    if REPORT.exists():
        errors = json.loads(REPORT.read_text()).get('errores')
        if errors: raise RuntimeError('Actualización parcial; se conservaron hojas fallidas: ' + '; '.join(errors))

if __name__ == '__main__': main()
