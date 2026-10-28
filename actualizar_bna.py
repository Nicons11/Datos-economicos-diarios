#!/usr/bin/env python3
"""Excel independiente. Nunca abre el archivo empresarial ni rutas de red."""
import argparse
import bisect
import hashlib
import json
import os
import re
import sys
from datetime import date, datetime, timedelta, timezone
from decimal import Decimal
from pathlib import Path
from zoneinfo import ZoneInfo

URL = 'https://www.errepar.com/cotizacion-dolar'
HEADERS = ['Fecha', 'Billete Compra', 'Billete Venta', 'Divisa Compra', 'Divisa Venta']
TZ = ZoneInfo('America/Argentina/Buenos_Aires')
ROOT = Path(__file__).resolve().parent

def parse_date(value):
    if isinstance(value, datetime):
        if value.time() != datetime.min.time():
            raise ValueError('La fecha contiene una hora')
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, str):
        for fmt in ('%d/%m/%Y', '%Y-%m-%d'):
            try:
                return datetime.strptime(value.strip(), fmt).date()
            except ValueError:
                pass
    raise ValueError('Usar fecha Excel, DD/MM/AAAA o AAAA-MM-DD; no números sueltos')

def parse_number(value):
    if not isinstance(value, str) or not re.fullmatch(r'(?:\d{1,3}(?:\.\d{3})+|\d+),\d{1,4}', value.strip()):
        raise ValueError(f'Formato de cotización inesperado: {value!r}')
    n = Decimal(value.strip().replace('.', '').replace(',', '.'))
    if not n.is_finite() or n <= 0:
        raise ValueError('Cotización no positiva')
    return n

def validate_snapshot(snap):
    if snap.get('url') != URL or snap.get('source') is not True or snap.get('headers') != HEADERS:
        raise ValueError('Fuente BNA o encabezados cambiaron; revisar captura')
    captured = datetime.fromisoformat(snap['captured_at'].replace('Z', '+00:00'))
    if captured.tzinfo is None:
        raise ValueError('Captura sin zona horaria')
    lo, hi = date.fromisoformat(snap['queried_from']), date.fromisoformat(snap['queried_to'])
    records = {}
    for row in snap['rows']:
        if len(row) != 5:
            raise ValueError('Fila de fuente incompleta')
        d = parse_date(row[0])
        if not lo <= d <= hi or d > captured.astimezone(TZ).date():
            raise ValueError('Fecha fuera de consulta o futura en la fuente')
        rates = tuple(parse_number(v) for v in row[1:])
        if rates[0] > rates[1] or rates[2] > rates[3]:
            raise ValueError('Compra mayor que venta; revisar fuente')
        if d in records:
            raise ValueError(f'Fecha repetida en fuente: {d}')
        records[d] = rates
    if not records:
        raise ValueError('Fuente sin filas: no se reemplazan resultados')
    return records, captured

def resolve(inputs, records, today, provisional_date=None):
    """Una salida por cada fila de entrada, sin ordenar ni deduplicar las entradas."""
    provisional_date = provisional_date or today
    keys = sorted(records)
    out = []
    for rownum, original in inputs:
        if original is None or original == '':
            out.append([None] * 8 + ['FILA VACÍA', rownum])
            continue
        try:
            d = parse_date(original)
        except ValueError as exc:
            out.append([str(original), None] + [None] * 6 + [f'FECHA INVÁLIDA: {exc}', rownum])
            continue
        if d > today:
            out.append([d, None] + [None] * 6 + ['FECHA FUTURA: sin valor', rownum])
            continue
        idx = bisect.bisect_right(keys, d) - 1
        if idx < 0:
            out.append([d, None] + [None] * 6 + ['SIN ANTECEDENTE EN CAPTURA: ampliar consulta', rownum])
            continue
        effective = keys[idx]
        lag = (d - effective).days
        state = 'EXACTA' if lag == 0 else 'ANTERIOR DISPONIBLE'
        if lag > 7:
            state += '; REVISAR: más de 7 días de arrastre'
        if effective == provisional_date:
            state += '; PROVISIONAL: revisar cierre'
        out.append([d, None, effective, lag, *[float(v) for v in records[effective]], state, rownum])
    return out

def local_path(value):
    # Solo archivos dentro de esta carpeta independiente; evita rutas privadas por accidente.
    p = Path(value)
    p = (ROOT / p).resolve() if not p.is_absolute() else p.resolve()
    if not p.is_relative_to(ROOT):
        raise ValueError('Usar un archivo dentro de la carpeta automatizacion_bna')
    return p

def fetch_live(first, last, directory):
    """Consulta el sitio renderizado, sin API privada ni cambio de proveedor."""
    from playwright.sync_api import sync_playwright
    directory.mkdir(parents=True, exist_ok=False)
    batches = []
    allrows = {}
    with sync_playwright() as pw:
        browser = pw.chromium.launch(headless=True, executable_path=os.environ.get('BNA_CHROMIUM_EXECUTABLE') or None)
        page = browser.new_page(locale='es-AR', timezone_id='America/Argentina/Buenos_Aires')
        try:
            page.goto(URL, wait_until='networkidle', timeout=60000)
            page.locator('table').wait_for(timeout=60000)
            # Bloques de hasta 31 días; retrocede hasta hallar antecedente o llegar al límite.
            end, floor = last, first - timedelta(days=3650)
            count = 0
            while end >= floor:
                start = max(end - timedelta(days=30), floor)
                page.locator('input[type=date]').nth(0).fill(start.isoformat())
                page.locator('input[type=date]').nth(1).fill(end.isoformat())
                page.get_by_text('VER DATOS', exact=True).click()
                page.wait_for_load_state('networkidle', timeout=60000)
                page.wait_for_timeout(1500)
                snap = page.locator('table').evaluate('''t => ({
                  headers: Array.from(t.querySelectorAll('th')).map(e=>e.textContent.trim()),
                  rows: Array.from(t.querySelectorAll('tr')).map(r=>Array.from(r.querySelectorAll('td')).map(e=>e.textContent.trim())).filter(r=>r.length>0 && !(r.length===1 && r[0].includes('Fuente: BNA'))),
                  source: t.innerText.includes('Fuente: BNA')})''')
                snap.update(url=URL, captured_at=datetime.now(timezone.utc).isoformat(),
                            queried_from=start.isoformat(), queried_to=end.isoformat(),
                            warning='El dato de la fecha de captura se trata como provisional')
                count += 1
                (directory / f'consulta_{count:03}.json').write_text(json.dumps(snap, ensure_ascii=False, indent=2))
                (directory / f'consulta_{count:03}.html').write_text(page.locator('table').evaluate('(t)=>t.outerHTML'))
                page.screenshot(path=str(directory / f'consulta_{count:03}.png'), full_page=True)
                # Vacío/cambio de estructura detiene el proceso, nunca se interpreta como feriado.
                records, _ = validate_snapshot(snap)
                for row in snap['rows']:
                    d = parse_date(row[0])
                    if d in allrows:
                        raise ValueError('Solapamiento inesperado de consultas')
                    allrows[d] = row
                batches.append(snap)
                if min(records) <= first:
                    break
                end = start - timedelta(days=1)
            if not allrows or min(allrows) > first:
                raise ValueError('No se encontró antecedente tras buscar hasta 10 años; revisar manualmente')
            merged = dict(batches[-1])
            merged.update(queried_from=min(date.fromisoformat(s['queried_from']) for s in batches).isoformat(),
                          queried_to=last.isoformat(), rows=[allrows[d] for d in sorted(allrows)],
                          captured_at=datetime.now(timezone.utc).isoformat())
            path = directory / 'fuente.json'
            path.write_text(json.dumps(merged, ensure_ascii=False, indent=2))
            return merged, path
        finally:
            browser.close()

def run(args):
    from openpyxl import load_workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    infile, outfile = local_path(args.archivo), local_path(args.salida)
    if infile == outfile or outfile.exists():
        raise ValueError('La salida debe ser nueva y distinta de la entrada; usar otro nombre')
    wb = load_workbook(infile)
    ws = wb['Fechas']
    lastrow = max((r for r in range(2, ws.max_row + 1) if ws.cell(r, 1).value not in (None, '')), default=1)
    if lastrow == 1:
        raise ValueError('Pegá primero las fechas en Fechas!A2, sin encabezado')
    if lastrow > 10001:
        raise ValueError('Máximo del prototipo: 10.000 filas de entrada')
    inputs = [(r, ws.cell(r, 1).value) for r in range(2, lastrow + 1)]
    for r, v in inputs:
        if ws.cell(r, 1).data_type == 'f':
            raise ValueError(f'Fechas!A{r}: pegar solo valores, sin fórmulas')
    selected = ws['D2'].value
    if selected not in ['PENDIENTE', *HEADERS[1:]]:
        raise ValueError('Selección inválida en Fechas!D2')
    today = datetime.now(TZ).date()
    valid = []
    for _, v in inputs:
        try:
            d = parse_date(v)
            if d <= today:
                valid.append(d)
        except ValueError:
            pass
    if not valid:
        raise ValueError('No hay fechas válidas no futuras para consultar')
    if args.snapshot:
        evidence = local_path(args.snapshot)
        snap = json.loads(evidence.read_text())
        mode = 'CAPTURA GUARDADA: no es una consulta actual'
    else:
        stamp = datetime.now(TZ).strftime('%Y%m%d_%H%M%S_%f')
        snap, evidence = fetch_live(min(valid), max(valid), ROOT / 'evidencia' / stamp)
        mode = 'CONSULTA WEB'
    records, captured = validate_snapshot(snap)
    if date.fromisoformat(snap['queried_to']) < max(valid):
        raise ValueError('La captura no cubre la fecha máxima de entrada; consultar de nuevo')
    provisional_date = captured.astimezone(TZ).date()
    results = resolve(inputs, records, today, provisional_date)
    # Aun una captura antigua permite reproducir historia, con procedencia visible.
    for row in results:
        if selected != 'PENDIENTE' and row[2] is not None:
            row[1] = row[4 + HEADERS[1:].index(selected)]
        if selected == 'PENDIENTE' and row[2] is not None:
            row[8] += '; FALTA DEFINIR SERIE'
    out = wb['Resultados']
    if out.max_row >= 7:
        out.delete_rows(7, out.max_row - 6)
    out['A3'] = f'Serie a copiar: {selected}. Valores fijos; volver a ejecutar al cambiar entradas o serie.'
    out['A4'] = f'{mode}. Captura: {captured.astimezone(TZ).isoformat(timespec="seconds")}'
    for r, row in enumerate(results, 7):
        for c, value in enumerate(row, 1):
            cell = out.cell(r, c, value)
            cell.font = Font(name='Arial', size=10)
            cell.alignment = Alignment(vertical='center')
            if c in (1, 3) and isinstance(value, date):
                cell.number_format = 'dd/mm/yyyy'
            if c in (2, 5, 6, 7, 8):
                cell.number_format = '#,##0.0000'
        if 'FALTA' in row[8] or 'PROVISIONAL' in row[8] or row[2] is None or row[3] > 7:
            out.cell(r, 9).fill = PatternFill('solid', fgColor='FFF0CD')
    out.auto_filter.ref = f'A6:J{6+len(results)}'
    source = wb['Fuente']
    source.delete_rows(1, source.max_row)
    meta = [
        ['Fuente', URL], ['Modo', mode], ['Captura (Argentina)', captured.astimezone(TZ).isoformat()],
        ['Intervalo consultado', f'{snap["queried_from"]} a {snap["queried_to"]}'],
        ['Evidencia', str(evidence.relative_to(ROOT))],
        ['SHA256 JSON', hashlib.sha256(evidence.read_bytes()).hexdigest()],
        ['Nota', 'Fecha efectiva = fecha de fila Errepar; no implica fecha de operación bancaria.'],
        ['Advertencia fuente', snap.get('warning', '')], [], HEADERS]
    for row in meta:
        source.append(row)
    for d in sorted(records):
        source.append([d, *[float(x) for x in records[d]]])
    for row in source.iter_rows(min_row=11):
        row[0].number_format = 'dd/mm/yyyy'
        for cell in row[1:5]:
            cell.number_format = '#,##0.0000'
    for cell in source[10]:
        cell.fill = PatternFill('solid', fgColor='23465D')
        cell.font = Font(name='Arial', bold=True, color='FFFFFF')
    outfile.parent.mkdir(parents=True, exist_ok=True)
    temporary = outfile.with_suffix('.tmp.xlsx')
    wb.save(temporary)
    temporary.replace(outfile)
    problems = sum(r[2] is None or selected == 'PENDIENTE' or r[2] == provisional_date or (r[3] or 0) > 7 for r in results)
    print(f'Creado: {outfile.name}. Filas: {len(results)}. Filas para revisar: {problems}. Serie: {selected}.')

if __name__ == '__main__':
    p = argparse.ArgumentParser(description=__doc__)
    p.add_argument('--archivo', default='BNA_independiente.xlsx')
    p.add_argument('--salida', default='BNA_actualizado.xlsx')
    p.add_argument('--snapshot', help='JSON de evidencia guardada, para reproducción sin Internet')
    args = p.parse_args()
    try:
        run(args)
    except Exception as exc:
        print(f'ERROR: {exc}. No se publica un nuevo Excel con resultados.', file=sys.stderr)
        sys.exit(1)
